"""
Aggregate Weighted ArcFace hygiene sweep results.

Analogous to reid_openset's `python -u -m utils.reid_plotting --aggregate`.

The shared plotting module expects results with a float 'threshold' config key
and filenames matching threshold=*_gallery=*_seed=*.json.  This script:

1. Loads arcface_weighted results and injects a synthetic threshold field
   (alpha=0 -> 0.0, alpha=1 -> 0.5, alpha=2 -> 1.0)
2. Writes patched copies with the expected filename pattern to a temp dir
3. Runs full per-backbone diagnostics (tables + diagnostic figures)
4. Runs cross-backbone aggregate analysis (faceted figures + tables)
5. Generates test performance tables with correct task-name mapping

Usage:
    cd /home/kdoherty/wolverines
    python -u arcface_weighted/aggregate_results.py
"""

import os
import sys
import json
import glob
import shutil
import tempfile
import numpy as np
from scipy import stats

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from utils.reid_config import MODEL_CONFIGS

ALPHA_TO_THRESHOLD = {
    0: 0.0,
    1: 0.5,
    2: 1.0,
}

BASE_DIR = "arcface_weighted"


# ============================================================================
# Test result loading (arcface_weighted paths + task names)
# ============================================================================

def load_test_results(model_name):
    """Load test evaluation results for a backbone from arcface_weighted paths.

    Returns:
        Dict mapping task name -> list of result dicts, or empty dict if none found.
    """
    config = MODEL_CONFIGS[model_name]
    experiment_dir = config['experiment_dir'].replace("reid_openset/", "arcface_weighted/", 1)
    pattern = os.path.join(experiment_dir, 'results', 'test', 'test_seed=*_*.json')
    json_files = glob.glob(pattern)

    if not json_files:
        return {}

    by_task = {}
    for json_file in json_files:
        try:
            with open(json_file, 'r') as f:
                data = json.load(f)
            task = data['config']['task']
            by_task.setdefault(task, []).append(data)
        except Exception as e:
            print(f"  Warning: failed to load {json_file}: {e}")

    return by_task


def create_test_performance_table(model_data, output_dir, metric_type):
    """Write test performance as an arXiv-style formatted .xlsx spreadsheet.

    Task mapping for arcface_weighted:
      closedset: baseline=closed_filtered (any alpha), weighted=closed_weighted (alpha>0)
      openset:   baseline=open_filtered (any alpha),   weighted=open_weighted (alpha>0)
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.utils import get_column_letter

    if metric_type == 'closedset':
        task_baseline = 'closed_filtered'
        task_weighted = 'closed_weighted'
        score_key = 'recall_at_1'
        score_label = 'R@1'
        filename = 'test_closedset.xlsx'
        sheet_title = 'Closed-Set Test'
    else:
        task_baseline = 'open_filtered'
        task_weighted = 'open_weighted'
        score_key = 'balanced_accuracy'
        score_label = 'BA'
        filename = 'test_openset.xlsx'
        sheet_title = 'Open-Set Test'

    models = list(model_data.keys())
    rows = []

    for model_name in models:
        label = model_data[model_name]['label'].replace('Frozen ', '')
        test_results = load_test_results(model_name)

        if not test_results:
            print(f"  No test results for {model_name}, skipping in {filename}")
            continue

        for strategy, task_key in [('Any alpha', task_baseline), ('Weighted only', task_weighted)]:
            task_data = test_results.get(task_key, [])
            if not task_data:
                continue

            values = [d['results'][score_key] for d in task_data]
            n = len(values)
            mean_val = float(np.mean(values))

            if n > 1:
                sem = float(stats.sem(values))
                ci_half = stats.t.ppf(0.975, n - 1) * sem
            else:
                ci_half = 0.0

            ci_lower = mean_val - ci_half
            ci_upper = mean_val + ci_half
            score_str = f"{mean_val:.2f}({ci_lower:.2f},{ci_upper:.2f})"

            cfg = task_data[0]['config']
            rows.append({
                'Backbone': label,
                'Alpha Filter': strategy,
                score_label: score_str,
                'Alpha': cfg.get('alpha', ''),
                'Query Thresh': cfg.get('query_quality_threshold', ''),
                'Epoch': cfg.get('best_epoch', ''),
            })

    if not rows:
        print(f"  No test results found for any backbone ({metric_type}), skipping table")
        return

    columns = ['Backbone', 'Alpha Filter', score_label, 'Alpha', 'Query Thresh', 'Epoch']

    serif_font = Font(name='Times New Roman', size=11)
    serif_bold = Font(name='Times New Roman', size=11, bold=True)
    header_fill = PatternFill(start_color='D9D9D9', end_color='D9D9D9', fill_type='solid')
    thin_side = Side(style='thin')
    thin_border = Border(left=thin_side, right=thin_side,
                         top=thin_side, bottom=thin_side)
    header_border = Border(left=thin_side, right=thin_side,
                           top=thin_side, bottom=Side(style='medium'))
    center_align = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center')

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title

    for col_idx, col_name in enumerate(columns, start=1):
        ws.cell(row=1, column=col_idx, value=col_name)

    for r_idx, row_data in enumerate(rows, start=2):
        for col_idx, col_name in enumerate(columns, start=1):
            ws.cell(row=r_idx, column=col_idx, value=row_data[col_name])

    total_cols = len(columns)
    for col_idx in range(1, total_cols + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = serif_bold
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = header_border
    ws.cell(row=1, column=1).alignment = left_align

    last_data_row = 1 + len(rows)
    for row_idx in range(2, last_data_row + 1):
        for col_idx in range(1, total_cols + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = serif_font
            cell.border = thin_border
            if col_idx <= 2:
                cell.alignment = left_align
            else:
                cell.alignment = center_align

    for col_idx in range(1, total_cols + 1):
        max_len = 0
        col_letter = get_column_letter(col_idx)
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx,
                                min_row=1, max_row=last_data_row):
            for cell in row:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, filename)
    wb.save(output_path)
    print(f"Saved test performance table: {output_path}")


# ============================================================================
# Main aggregate pipeline
# ============================================================================

def run_single_model(model_name, tmp_results_dir):
    """Per-backbone diagnostics: tables + diagnostic figures.

    Returns dict with strategies + results for cross-backbone aggregation,
    or None if no results found.
    """
    from utils.reid_plotting import (
        load_hygiene_results, print_summary_table,
        collect_strategies_data, save_scores_table,
        save_thresholds_table, save_summary_table,
        plot_recall_curves, plot_validation_loss_curves,
        plot_cosine_threshold,
    )

    print("=" * 60)
    print(f"Weighted ArcFace Evaluation: {model_name} Results Analysis")
    print("=" * 60)

    results = load_hygiene_results(tmp_results_dir)
    if len(results) == 0:
        print(f"No results found for {model_name}. Skipping.")
        return None

    print_summary_table(results)

    # Tables -> arcface_weighted/tables/{model_name}/
    tables_dir = os.path.join(BASE_DIR, 'tables', model_name)
    os.makedirs(tables_dir, exist_ok=True)

    strategies = collect_strategies_data(results)
    save_scores_table(strategies, tables_dir)
    save_thresholds_table(strategies, tables_dir)
    save_summary_table(results, tables_dir)

    # Diagnostic figures -> arcface_weighted/figures/{model_name}/
    figures_dir = os.path.join(BASE_DIR, 'figures', model_name)
    os.makedirs(figures_dir, exist_ok=True)

    plot_recall_curves(results, os.path.join(figures_dir, 'recall_curves.png'))
    plot_validation_loss_curves(results, os.path.join(figures_dir, 'validation_loss_curves.png'))
    plot_cosine_threshold(results, os.path.join(figures_dir, 'cosine_threshold.png'))

    print(f"\nDiagnostics complete for {model_name}!")
    print(f"  Tables:  {tables_dir}")
    print(f"  Figures: {figures_dir}")

    return {'strategies': strategies, 'results': results}


def patch_hygiene_results(model_name):
    """Load arcface_weighted hygiene results, inject synthetic threshold, write to temp dir.

    Returns (tmp_dir, n_files) or (None, 0) if no results found.
    """
    config = MODEL_CONFIGS[model_name]
    experiment_dir = config['experiment_dir'].replace("reid_openset/", "arcface_weighted/", 1)
    src_dir = os.path.join(experiment_dir, 'results', 'hygiene')

    src_files = glob.glob(os.path.join(src_dir, "alpha=*_gallery=*_seed=*.json"))
    if not src_files:
        print(f"No arcface_weighted hygiene results found in {src_dir}")
        return None, 0

    print(f"Found {len(src_files)} hygiene result files for {model_name}")

    tmp_dir = tempfile.mkdtemp(prefix=f"arcface_weighted_{model_name}_")

    for src_file in src_files:
        with open(src_file, 'r') as f:
            data = json.load(f)

        alpha = data['config']['alpha']
        threshold = ALPHA_TO_THRESHOLD[alpha]
        gallery_size = data['config']['gallery_size']
        seed = data['config']['seed']

        data['config']['threshold'] = threshold

        dst_name = f"threshold={threshold:.2f}_gallery={gallery_size}_seed={seed}.json"
        dst_path = os.path.join(tmp_dir, dst_name)

        with open(dst_path, 'w') as f:
            json.dump(data, f, indent=2)

    print(f"Wrote {len(src_files)} patched files to {tmp_dir}")
    return tmp_dir, len(src_files)


def main():
    from utils.reid_plotting import (
        create_rank1_figure, create_novelty_detection_figure,
        create_rank1_thresholds_table, create_novelty_thresholds_table,
    )

    print("=" * 60)
    print("Weighted ArcFace Cross-Backbone Aggregate Analysis")
    print("=" * 60)

    model_data = {}
    tmp_dirs = []

    for model_name, config in MODEL_CONFIGS.items():
        print(f"\n--- Loading {model_name} ---")

        tmp_dir, n_files = patch_hygiene_results(model_name)
        if tmp_dir is None:
            print(f"Skipping {model_name} (no results)")
            continue
        tmp_dirs.append(tmp_dir)

        result = run_single_model(model_name, tmp_dir)
        if result is None:
            continue

        model_data[model_name] = {
            'strategies': result['strategies'],
            'results': result['results'],
            'label': config['backbone_label'],
        }

    if not model_data:
        print("No backbone results found. Exiting.")
        for d in tmp_dirs:
            shutil.rmtree(d, ignore_errors=True)
        return

    try:
        # Cross-backbone figures -> arcface_weighted/figures/
        cross_figures_dir = os.path.join(BASE_DIR, 'figures')
        os.makedirs(cross_figures_dir, exist_ok=True)

        # Cross-backbone tables -> arcface_weighted/tables/
        cross_tables_dir = os.path.join(BASE_DIR, 'tables')
        os.makedirs(cross_tables_dir, exist_ok=True)

        print(f"\nGenerating cross-backbone figures ({len(model_data)} backbones)...")
        create_rank1_figure(model_data, cross_figures_dir)
        create_novelty_detection_figure(model_data, cross_figures_dir)
        create_rank1_thresholds_table(model_data, cross_tables_dir)
        create_novelty_thresholds_table(model_data, cross_tables_dir)

        # Test performance tables (arcface_weighted task names)
        print(f"\nGenerating test performance tables...")
        create_test_performance_table(model_data, cross_tables_dir, 'closedset')
        create_test_performance_table(model_data, cross_tables_dir, 'openset')

        print(f"\nAggregate analysis complete!")
        print(f"  Cross-backbone figures: {cross_figures_dir}")
        print(f"  Cross-backbone tables:  {cross_tables_dir}")

    finally:
        for d in tmp_dirs:
            shutil.rmtree(d, ignore_errors=True)


if __name__ == "__main__":
    main()
