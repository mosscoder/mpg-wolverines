"""
Table generation functions for reid plotting: scores, thresholds, test performance.
"""

import os
import json
import glob
import numpy as np
import pandas as pd
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from scipy import stats


def save_scores_table(strategies: dict, output_dir: str):
    """Save scores data as CSV tables."""
    # R@1 comparison table
    r1_data = []
    for i, x in enumerate(strategies['baseline']['x']):
        row = {'gallery_size': x}
        if i < len(strategies['baseline']['r1']):
            row['baseline_r1'] = strategies['baseline']['r1'][i]
            row['baseline_r1_ci_lower'] = strategies['baseline']['r1_ci_lower'][i]
            row['baseline_r1_ci_upper'] = strategies['baseline']['r1_ci_upper'][i]
        r1_data.append(row)

    for i, x in enumerate(strategies['optimal']['x']):
        match = next((r for r in r1_data if r['gallery_size'] == x), None)
        if match is None:
            match = {'gallery_size': x}
            r1_data.append(match)
        if i < len(strategies['optimal']['r1']):
            match['optimal_r1'] = strategies['optimal']['r1'][i]
            match['optimal_r1_ci_lower'] = strategies['optimal']['r1_ci_lower'][i]
            match['optimal_r1_ci_upper'] = strategies['optimal']['r1_ci_upper'][i]
            if i < len(strategies['optimal']['best_gal']):
                match['optimal_gal_thresh'] = strategies['optimal']['best_gal'][i]
            if i < len(strategies['optimal']['best_q']):
                match['optimal_query_thresh'] = strategies['optimal']['best_q'][i]

    r1_df = pd.DataFrame(r1_data).sort_values('gallery_size')
    if 'baseline_r1' in r1_df.columns and 'optimal_r1' in r1_df.columns:
        r1_df['delta_r1'] = r1_df['optimal_r1'] - r1_df['baseline_r1']

    r1_path = os.path.join(output_dir, 'scores_r1.csv')
    r1_df.to_csv(r1_path, index=False, float_format='%.4f')
    print(f"Saved R@1 scores table: {r1_path}")

    # Balanced Accuracy comparison table
    ba_data = []
    for i, x in enumerate(strategies['baseline']['x']):
        row = {'gallery_size': x}
        if i < len(strategies['baseline']['ba']):
            row['baseline_ba'] = strategies['baseline']['ba'][i]
            row['baseline_ba_ci_lower'] = strategies['baseline']['ba_ci_lower'][i]
            row['baseline_ba_ci_upper'] = strategies['baseline']['ba_ci_upper'][i]
        ba_data.append(row)

    for i, x in enumerate(strategies['optimal_ba']['x']):
        match = next((r for r in ba_data if r['gallery_size'] == x), None)
        if match is None:
            match = {'gallery_size': x}
            ba_data.append(match)
        if i < len(strategies['optimal_ba']['ba']):
            match['optimal_ba'] = strategies['optimal_ba']['ba'][i]
            match['optimal_ba_ci_lower'] = strategies['optimal_ba']['ba_ci_lower'][i]
            match['optimal_ba_ci_upper'] = strategies['optimal_ba']['ba_ci_upper'][i]
            if i < len(strategies['optimal_ba']['best_gal']):
                match['optimal_gal_thresh'] = strategies['optimal_ba']['best_gal'][i]
            if i < len(strategies['optimal_ba']['best_q']):
                match['optimal_query_thresh'] = strategies['optimal_ba']['best_q'][i]

    ba_df = pd.DataFrame(ba_data).sort_values('gallery_size')
    if 'baseline_ba' in ba_df.columns and 'optimal_ba' in ba_df.columns:
        ba_df['delta_ba'] = ba_df['optimal_ba'] - ba_df['baseline_ba']

    ba_path = os.path.join(output_dir, 'scores_ba.csv')
    ba_df.to_csv(ba_path, index=False, float_format='%.4f')
    print(f"Saved BA scores table: {ba_path}")


def save_thresholds_table(strategies: dict, output_dir: str):
    """Save optimal thresholds data as CSV tables."""
    r1_thresh_data = []
    for i, x in enumerate(strategies['optimal']['x']):
        row = {'gallery_size': x}
        if i < len(strategies['optimal']['best_gal']):
            row['gallery_threshold'] = strategies['optimal']['best_gal'][i]
        if i < len(strategies['optimal']['best_q']):
            row['query_threshold'] = strategies['optimal']['best_q'][i]
        r1_thresh_data.append(row)

    r1_df = pd.DataFrame(r1_thresh_data).sort_values('gallery_size')
    r1_path = os.path.join(output_dir, 'thresholds_r1_optimized.csv')
    r1_df.to_csv(r1_path, index=False, float_format='%.2f')
    print(f"Saved R@1-optimized thresholds table: {r1_path}")

    ba_thresh_data = []
    for i, x in enumerate(strategies['optimal_ba']['x']):
        row = {'gallery_size': x}
        if i < len(strategies['optimal_ba']['best_gal']):
            row['gallery_threshold'] = strategies['optimal_ba']['best_gal'][i]
        if i < len(strategies['optimal_ba']['best_q']):
            row['query_threshold'] = strategies['optimal_ba']['best_q'][i]
        ba_thresh_data.append(row)

    ba_df = pd.DataFrame(ba_thresh_data).sort_values('gallery_size')
    ba_path = os.path.join(output_dir, 'thresholds_ba_optimized.csv')
    ba_df.to_csv(ba_path, index=False, float_format='%.2f')
    print(f"Saved BA-optimized thresholds table: {ba_path}")


def load_test_results(model_name: str) -> dict:
    """Load test evaluation results for a backbone, grouped by task.

    Args:
        model_name: Key in MODEL_CONFIGS (e.g. 'dinov3')

    Returns:
        Dict mapping task name -> list of result dicts, or empty dict if none found.
    """
    from utils.reid_config import MODEL_CONFIGS
    config = MODEL_CONFIGS[model_name]
    pattern = os.path.join(config['experiment_dir'], 'results', 'test', 'test_seed=*_*.json')
    json_files = glob.glob(pattern)

    if not json_files:
        return {}

    by_task = {}
    for json_file in json_files:
        try:
            with open(json_file, 'r') as f:
                data = json.load(f)
            task = data['config']['task']
            by_task.setdefault(task, []).append(data)
        except Exception as e:
            print(f"  Warning: failed to load {json_file}: {e}")

    return by_task


def create_thresholds_table(model_data: dict, strategy_key: str,
                            output_dir: str, filename: str):
    """Write optimal thresholds as an arXiv-style formatted .xlsx spreadsheet.

    Layout (example with 3 backbones):

        Examples per   DINOv3-ViT-B/16    MegaDescriptor-L-384   BioCLIP-2 ViT-L/14
        individual    Training Validation  Training Validation    Training Validation
        2             0.20     0.50        ...      ...           ...      ...
        4             0.30     0.50        ...      ...           ...      ...
        ...

    Args:
        model_data: Dict mapping model name -> {"strategies", "results", "label"}
        strategy_key: 'optimal' (R@1-optimized) or 'optimal_ba' (BA-optimized)
        output_dir: Directory to save xlsx
        filename: Output filename (e.g. 'best_threshold_rank.xlsx')
    """
    from openpyxl import Workbook

    models = list(model_data.keys())
    labels = [model_data[m]['label'].replace('Frozen ', '') for m in models]

    # Collect all gallery sizes across backbones
    all_sizes = sorted(set(
        sz for m in models
        for sz in model_data[m]['strategies'][strategy_key]['x']
    ))

    # Build lookup: model -> gallery_size -> (gal_thresh, q_thresh)
    lookup = {}
    for m in models:
        s = model_data[m]['strategies'][strategy_key]
        lookup[m] = {}
        for i, sz in enumerate(s['x']):
            gal = s['best_gal'][i] if i < len(s['best_gal']) else None
            q = s['best_q'][i] if i < len(s['best_q']) else None
            lookup[m][sz] = (gal, q)

    # Style definitions
    serif_font = Font(name='Times New Roman', size=11)
    serif_bold = Font(name='Times New Roman', size=11, bold=True)
    header_fill = PatternFill(start_color='D9D9D9', end_color='D9D9D9', fill_type='solid')
    thin_side = Side(style='thin')
    thick_side = Side(style='medium')
    thin_border = Border(left=thin_side, right=thin_side,
                         top=thin_side, bottom=thin_side)
    header_border = Border(left=thin_side, right=thin_side,
                           top=thin_side, bottom=thick_side)
    center_align = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center')

    n_data_cols = len(models) * 2
    total_cols = 1 + n_data_cols

    wb = Workbook()
    ws = wb.active
    ws.title = 'Thresholds'

    # Row 1: backbone names (merged across Training+Validation pairs)
    ws.cell(row=1, column=1)  # empty corner cell
    for k, label in enumerate(labels):
        col = 2 + k * 2
        ws.cell(row=1, column=col, value=label)
        # Write to both cells BEFORE merging so formatting sticks
        ws.cell(row=1, column=col + 1, value='')

    # Row 2: "Examples per individual" + Training/Validation sub-headers
    ws.cell(row=2, column=1, value='Examples per individual')
    for k in range(len(models)):
        ws.cell(row=2, column=2 + k * 2, value='Training')
        ws.cell(row=2, column=2 + k * 2 + 1, value='Validation')

    # Data rows (row 3+)
    for r_idx, sz in enumerate(all_sizes):
        row = 3 + r_idx
        ws.cell(row=row, column=1, value=sz)
        for k, m in enumerate(models):
            col = 2 + k * 2
            if sz in lookup[m] and lookup[m][sz][0] is not None:
                gal, q = lookup[m][sz]
                ws.cell(row=row, column=col, value=gal)
                ws.cell(row=row, column=col + 1, value=q)

    # Apply formatting BEFORE merging (merged cells lose right-cell styles)
    # Header row 1
    for col_idx in range(1, total_cols + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = serif_bold
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    ws.cell(row=1, column=1).alignment = left_align

    # Sub-header row 2
    for col_idx in range(1, total_cols + 1):
        cell = ws.cell(row=2, column=col_idx)
        cell.font = serif_bold
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = header_border
    ws.cell(row=2, column=1).alignment = left_align

    # Data rows
    last_data_row = 2 + len(all_sizes)
    for row_idx in range(3, last_data_row + 1):
        for col_idx in range(1, total_cols + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = serif_font
            cell.border = thin_border
            if col_idx == 1:
                cell.alignment = left_align
                cell.font = serif_bold
            else:
                cell.alignment = center_align
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '0.00'

    # Now merge backbone name cells (after formatting is applied)
    for k in range(len(labels)):
        start_col = 2 + k * 2
        ws.merge_cells(start_row=1, start_column=start_col,
                       end_row=1, end_column=start_col + 1)

    # Auto-fit column widths
    for col_idx in range(1, total_cols + 1):
        max_len = 0
        col_letter = get_column_letter(col_idx)
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx,
                                min_row=1, max_row=last_data_row):
            for cell in row:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, filename)
    wb.save(output_path)
    print(f"Saved thresholds table: {output_path}")


def create_test_performance_table(model_data: dict, output_dir: str,
                                  metric_type: str):
    """Write test performance as an arXiv-style formatted .xlsx spreadsheet.

    Produces one table with two rows per backbone (None + Optimal filter strategy),
    showing mean(95% CI) scores across seeds.

    Args:
        model_data: Dict mapping model name -> {"strategies", "results", "label"}
        output_dir: Directory to save xlsx
        metric_type: 'closedset' or 'openset'
    """
    from openpyxl import Workbook

    if metric_type == 'closedset':
        task_none = 'closed_unfiltered'
        task_optimal = 'closed_filtered'
        score_key = 'recall_at_1'
        score_label = 'R@1'
        filename = 'test_closedset.xlsx'
        sheet_title = 'Closed-Set Test'
    else:
        task_none = 'open_unfiltered'
        task_optimal = 'open_filtered'
        score_key = 'balanced_accuracy'
        score_label = 'BA'
        filename = 'test_openset.xlsx'
        sheet_title = 'Open-Set Test'

    models = list(model_data.keys())
    rows = []

    for model_name in models:
        label = model_data[model_name]['label'].replace('Frozen ', '')
        test_results = load_test_results(model_name)

        if not test_results:
            print(f"  No test results for {model_name}, skipping in {filename}")
            continue

        for strategy, task_key in [('None', task_none), ('Optimal', task_optimal)]:
            task_data = test_results.get(task_key, [])
            if not task_data:
                continue

            values = [d['results'][score_key] for d in task_data]
            n = len(values)
            mean_val = float(np.mean(values))

            if n > 1:
                sem = float(stats.sem(values))
                ci_half = stats.t.ppf(0.975, n - 1) * sem
            else:
                ci_half = 0.0

            ci_lower = mean_val - ci_half
            ci_upper = mean_val + ci_half
            score_str = f"{mean_val:.2f}({ci_lower:.2f},{ci_upper:.2f})"

            # Extract hyperparams (same across seeds for a given task)
            cfg = task_data[0]['config']
            rows.append({
                'Backbone': label,
                'Filter Strategy': strategy,
                score_label: score_str,
                'Gallery Size': cfg.get('gallery_size', ''),
                'Gallery Thresh': cfg.get('threshold', ''),
                'Query Thresh': cfg.get('query_quality_threshold', ''),
                'Epoch': cfg.get('best_epoch', ''),
            })

    if not rows:
        print(f"  No test results found for any backbone ({metric_type}), skipping table")
        return

    # Build xlsx with same styling as create_thresholds_table
    columns = ['Backbone', 'Filter Strategy', score_label,
               'Gallery Size', 'Gallery Thresh', 'Query Thresh', 'Epoch']

    serif_font = Font(name='Times New Roman', size=11)
    serif_bold = Font(name='Times New Roman', size=11, bold=True)
    header_fill = PatternFill(start_color='D9D9D9', end_color='D9D9D9', fill_type='solid')
    thin_side = Side(style='thin')
    thick_side = Side(style='medium')
    thin_border = Border(left=thin_side, right=thin_side,
                         top=thin_side, bottom=thin_side)
    header_border = Border(left=thin_side, right=thin_side,
                           top=thin_side, bottom=thick_side)
    center_align = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center')

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title

    # Header row
    for col_idx, col_name in enumerate(columns, start=1):
        ws.cell(row=1, column=col_idx, value=col_name)

    # Data rows
    for r_idx, row_data in enumerate(rows, start=2):
        for col_idx, col_name in enumerate(columns, start=1):
            ws.cell(row=r_idx, column=col_idx, value=row_data[col_name])

    # Format header
    total_cols = len(columns)
    for col_idx in range(1, total_cols + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = serif_bold
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = header_border
    ws.cell(row=1, column=1).alignment = left_align

    # Format data rows
    last_data_row = 1 + len(rows)
    for row_idx in range(2, last_data_row + 1):
        for col_idx in range(1, total_cols + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = serif_font
            cell.border = thin_border
            if col_idx <= 2:
                cell.alignment = left_align
            else:
                cell.alignment = center_align

    # Auto-fit column widths
    for col_idx in range(1, total_cols + 1):
        max_len = 0
        col_letter = get_column_letter(col_idx)
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx,
                                min_row=1, max_row=last_data_row):
            for cell in row:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, filename)
    wb.save(output_path)
    print(f"Saved test performance table: {output_path}")


def create_rank1_thresholds_table(model_data: dict, output_dir: str):
    """Create cross-backbone table of R@1-optimized quality thresholds."""
    create_thresholds_table(
        model_data, 'optimal', output_dir,
        'best_threshold_rank.xlsx',
    )


def create_novelty_thresholds_table(model_data: dict, output_dir: str):
    """Create cross-backbone table of BA-optimized quality thresholds."""
    create_thresholds_table(
        model_data, 'optimal_ba', output_dir,
        'best_threshold_novelty.xlsx',
    )
